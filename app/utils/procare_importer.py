import io
import csv
import openpyxl
from datetime import datetime, date
from collections import defaultdict
from app.models import db
from app.models.student_model import Student, Parent
from app.models.financial_model import (
    StudentFinancialAccount,
    BillingPlan,
    Subscription,
    PresetChargeItem,
    PresetDiscount,
    FinancialAuditLog
)
from app.models.activity_log_model import log_activity
from dateutil.relativedelta import relativedelta


def parse_procare_date(d_val):
    if not d_val:
        return None
    if isinstance(d_val, datetime):
        return d_val.date()
    if isinstance(d_val, date):
        return d_val
    s = str(d_val).strip()
    for fmt in ('%d %B, %Y', '%Y-%m-%d', '%m/%d/%Y', '%B %d, %Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def room_to_grade(room):
    if not room:
        return "Unassigned"
    r = room.strip()
    if "Kindergarten" in r:
        return "Kindergarten"
    for i in range(1, 13):
        suffix = "th" if 4 <= i <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(i % 10, "th")
        if f"{i}{suffix}" in r or f"Room {i}" in r or f"Grade {i}" in r:
            return f"{i}{suffix} Grade"
    cleaned = r.replace("Home Room ", "").strip()
    return cleaned if cleaned else "Unassigned"


def parse_procare_rows(rows_iterator):
    """
    Parses rows from either an Excel worksheet or CSV reader.
    Finds header row containing 'Student Name', 'Plan Name', 'Item', 'Amount'.
    """
    header_idx = None
    headers = []
    data_rows = []

    for idx, row in enumerate(rows_iterator):
        row_vals = [str(c).strip() if c is not None else "" for c in row]
        # Look for the header row
        joined_lower = " ".join(row_vals).lower()
        if "student name" in joined_lower and ("plan name" in joined_lower or "item" in joined_lower):
            header_idx = idx
            headers = [h.strip() for h in row_vals]
            continue
        
        if header_idx is not None and any(row_vals):
            data_rows.append(row)

    if header_idx is None:
        raise ValueError("Could not find Procare header row containing 'Student Name', 'Plan Name', etc.")

    # Build column map
    col_map = {}
    for i, h in enumerate(headers):
        hl = h.lower()
        if "student name" in hl:
            col_map['student_name'] = i
        elif "room" in hl:
            col_map['room'] = i
        elif "tags" in hl:
            col_map['tags'] = i
        elif "family id" in hl:
            col_map['family_id'] = i
        elif "parent name" in hl:
            col_map['parents'] = i
        elif "plan name" in hl:
            col_map['plan_name'] = i
        elif "plan status" in hl:
            col_map['plan_status'] = i
        elif "plan cycle" in hl:
            col_map['cycle'] = i
        elif "plan start" in hl:
            col_map['start_date'] = i
        elif "plan end" in hl:
            col_map['end_date'] = i
        elif hl == "item" or "item" in hl:
            col_map['item'] = i
        elif hl == "amount":
            col_map['amount'] = i
        elif "total amount" in hl:
            col_map['total_amount'] = i

    students_dict = {}
    preset_charges = {}
    preset_discounts = set()
    templates_items = defaultdict(list)

    for row in data_rows:
        def get_val(key):
            idx = col_map.get(key)
            if idx is not None and idx < len(row):
                val = row[idx]
                return val if val is not None else ""
            return ""

        s_raw_name = str(get_val('student_name')).strip()
        if not s_raw_name or s_raw_name.lower() == 'student name':
            continue

        room = str(get_val('room')).strip()
        tags = str(get_val('tags')).strip()
        fam_id = str(get_val('family_id')).strip()
        parents = str(get_val('parents')).strip()
        plan_name = str(get_val('plan_name')).strip()
        plan_status = str(get_val('plan_status')).strip() or "Active"
        cycle = str(get_val('cycle')).strip() or "Monthly"
        start_date = parse_procare_date(get_val('start_date'))
        end_date = parse_procare_date(get_val('end_date'))
        item_desc = str(get_val('item')).strip()
        
        raw_amt = get_val('amount')
        amount = None
        if raw_amt != "" and raw_amt is not None:
            try:
                # Remove currency symbols or commas if in string
                if isinstance(raw_amt, str):
                    cleaned_amt = raw_amt.replace('$', '').replace(',', '').strip()
                    amount = float(cleaned_amt)
                else:
                    amount = float(raw_amt)
            except ValueError:
                amount = None

        raw_total = get_val('total_amount')
        total_amt = None
        if raw_total != "" and raw_total is not None:
            try:
                if isinstance(raw_total, str):
                    cleaned_tot = raw_total.replace('$', '').replace(',', '').strip()
                    total_amt = float(cleaned_tot)
                else:
                    total_amt = float(raw_total)
            except ValueError:
                total_amt = None

        if s_raw_name not in students_dict:
            if ',' in s_raw_name:
                parts = s_raw_name.split(',', 1)
                last_name = parts[0].strip()
                first_name = parts[1].strip()
            else:
                parts = s_raw_name.split()
                first_name = parts[0].strip()
                last_name = " ".join(parts[1:]).strip() if len(parts) > 1 else ""

            students_dict[s_raw_name] = {
                'raw_name': s_raw_name,
                'first_name': first_name,
                'last_name': last_name,
                'room': room,
                'grade_level': room_to_grade(room),
                'tags': tags,
                'family_id': fam_id,
                'parents': parents,
                'has_plan': bool(plan_name),
                'plan_name': plan_name if plan_name else None,
                'plan_status': plan_status,
                'cycle': cycle,
                'start_date': start_date.isoformat() if start_date else None,
                'end_date': end_date.isoformat() if end_date else None,
                'header_total': total_amt,
                'items': []
            }

        if item_desc and amount is not None:
            is_discount = amount < 0
            item_obj = {
                'description': item_desc,
                'amount': amount,
                'type': 'Discount' if is_discount else 'New Item',
                'value': abs(amount),
                'unit': '$'
            }
            students_dict[s_raw_name]['items'].append(item_obj)

            if is_discount:
                preset_discounts.add(item_desc)
            else:
                # Save preset charge with its amount
                if item_desc not in preset_charges or preset_charges[item_desc] == 0:
                    preset_charges[item_desc] = amount

            if plan_name:
                templates_items[plan_name].append(item_obj)

    # Calculate net totals for students
    for s_info in students_dict.values():
        calc_total = sum(i['amount'] for i in s_info['items'])
        s_info['calculated_total'] = round(calc_total, 2)
        if s_info['header_total'] is None:
            s_info['header_total'] = s_info['calculated_total']

    # Deduplicate template items to produce canonical template structures
    templates = []
    for p_name, t_items in templates_items.items():
        unique_map = {}
        for it in t_items:
            desc = it['description']
            if desc not in unique_map:
                unique_map[desc] = it
        templates.append({
            'name': p_name,
            'items_json': list(unique_map.values())
        })

    return {
        'students': list(students_dict.values()),
        'templates': templates,
        'preset_charges': [{'description': k, 'amount': v} for k, v in preset_charges.items()],
        'preset_discounts': [{'description': d} for d in preset_discounts]
    }


def parse_procare_excel_file(file_content_or_path):
    """Parses .xlsx file from binary stream or filepath."""
    if isinstance(file_content_or_path, (str, bytes, io.BytesIO)):
        if isinstance(file_content_or_path, str):
            wb = openpyxl.load_workbook(file_content_or_path, data_only=True)
        else:
            wb = openpyxl.load_workbook(io.BytesIO(file_content_or_path), data_only=True)
    else:
        wb = openpyxl.load_workbook(file_content_or_path, data_only=True)

    ws = wb.active
    rows = []
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        rows.append(vals)

    return parse_procare_rows(rows)


def parse_procare_csv_file(file_content_or_path):
    """Parses .csv file from string/bytes or filepath."""
    if isinstance(file_content_or_path, str) and not file_content_or_path.endswith('.csv'):
        # String content
        f = io.StringIO(file_content_or_path)
    elif isinstance(file_content_or_path, bytes):
        f = io.StringIO(file_content_or_path.decode('utf-8', errors='replace'))
    else:
        f = open(file_content_or_path, 'r', encoding='utf-8', errors='replace')

    reader = csv.reader(f)
    rows = list(reader)
    if hasattr(f, 'close') and f != file_content_or_path:
        f.close()
    return parse_procare_rows(rows)


def load_and_parse_file(file_obj, filename=""):
    """Auto-detects format (Excel or CSV) and parses."""
    fn = filename.lower()
    if fn.endswith('.csv') or fn.endswith('.txt'):
        content = file_obj.read()
        return parse_procare_csv_file(content)
    else:
        content = file_obj.read()
        return parse_procare_excel_file(content)


def execute_procare_import(parsed_data, options=None, actor=None):
    """
    Executes import of parsed Procare data into database.
    options:
      - import_templates: bool (default True)
      - import_presets: bool (default True)
      - create_missing_students: bool (default True)
      - import_subscriptions: bool (default True)
      - all_students: bool (default False - if True, creates students without plans too)
    """
    if options is None:
        options = {}

    import_templates = options.get('import_templates', True)
    import_presets = options.get('import_presets', True)
    create_missing_students = options.get('create_missing_students', True)
    import_subscriptions = options.get('import_subscriptions', True)
    import_all_students = options.get('all_students', False)

    today = date.today()
    results = {
        'templates_created': 0,
        'templates_updated': 0,
        'presets_created': 0,
        'discounts_created': 0,
        'students_matched': 0,
        'students_created': 0,
        'subscriptions_created': 0,
        'subscriptions_updated': 0,
        'skipped_no_plan': 0,
        'errors': []
    }

    # 1. Billing Plan Templates
    if import_templates:
        for tmpl in parsed_data.get('templates', []):
            name = tmpl['name']
            items_json = tmpl['items_json']
            existing = BillingPlan.query.filter_by(name=name).first()
            if existing:
                existing.items_json = items_json
                existing.is_active = True
                results['templates_updated'] += 1
            else:
                new_plan = BillingPlan(name=name, items_json=items_json, is_active=True)
                db.session.add(new_plan)
                results['templates_created'] += 1

    # 2. Preset Charges & Discounts
    if import_presets:
        for ch in parsed_data.get('preset_charges', []):
            desc = ch['description']
            amt = ch['amount']
            existing = PresetChargeItem.query.filter_by(description=desc).first()
            if not existing:
                db.session.add(PresetChargeItem(description=desc, amount=amt, is_active=True))
                results['presets_created'] += 1

        for d in parsed_data.get('preset_discounts', []):
            desc = d['description']
            existing = PresetDiscount.query.filter_by(description=desc).first()
            if not existing:
                db.session.add(PresetDiscount(description=desc, is_active=True))
                results['discounts_created'] += 1

    db.session.flush()

    # Pre-index existing students for fast matching
    all_db_students = Student.query.all()
    student_lookup = {}
    for st in all_db_students:
        key = (st.first_name.strip().lower(), st.last_name.strip().lower())
        student_lookup[key] = st

    # 3. Students & Subscriptions
    students_to_process = parsed_data.get('students', [])
    for s_info in students_to_process:
        has_plan = s_info.get('has_plan', False)
        if not has_plan and not import_all_students:
            results['skipped_no_plan'] += 1
            continue

        fn = s_info.get('first_name', '').strip()
        ln = s_info.get('last_name', '').strip()
        key = (fn.lower(), ln.lower())

        student = student_lookup.get(key)
        if not student:
            # Also try matching without middle names or by last name + first word of first name
            fn_first_word = fn.split()[0] if fn else ""
            fallback_key = (fn_first_word.lower(), ln.lower())
            student = student_lookup.get(fallback_key)

        if student:
            results['students_matched'] += 1
        else:
            if not create_missing_students:
                continue

            # Create Student
            room = s_info.get('room', '')
            grade = s_info.get('grade_level', room_to_grade(room))
            fam_id = s_info.get('family_id', '')
            p_names = s_info.get('parents', '')

            # Approximate birthdate based on grade level
            dob = date(2015, 1, 1)
            try:
                if 'kindergarten' in grade.lower():
                    dob = date(2020, 1, 1)
                else:
                    for g in range(1, 13):
                        if f"{g}" in grade:
                            dob = date(2020 - g, 1, 1)
                            break
            except Exception:
                pass

            s_date = parse_procare_date(s_info.get('start_date')) or today

            student = Student(
                first_name=fn,
                last_name=ln,
                grade_level=grade,
                date_of_birth=dob,
                status='Active',
                enrollment_date=s_date,
                notes=f"Imported from Procare. Family ID: {fam_id}, Room: {room}. Parents: {p_names}"
            )
            db.session.add(student)
            db.session.flush()

            # Create financial account
            fin_account = StudentFinancialAccount(student_id=student.id)
            db.session.add(fin_account)
            db.session.flush()

            # Handle Parents
            if p_names:
                p_list = [p.strip() for p in p_names.split(',') if p.strip()]
                # Deduplicate parent names
                p_list_unique = []
                for p in p_list:
                    if p not in p_list_unique:
                        p_list_unique.append(p)

                for p_full in p_list_unique:
                    p_parts = p_full.split()
                    p_fn = p_parts[0] if p_parts else "Parent"
                    p_ln = " ".join(p_parts[1:]) if len(p_parts) > 1 else ln
                    clean_email = f"{p_fn.lower()}.{p_ln.lower().replace(' ', '')}.{fam_id.lower() or student.id}@parent.elaaschool.org"

                    existing_parent = Parent.query.filter_by(email=clean_email).first()
                    if not existing_parent:
                        existing_parent = Parent(
                            first_name=p_fn,
                            last_name=p_ln,
                            email=clean_email,
                            phone="555-0100",
                            is_active=True
                        )
                        db.session.add(existing_parent)
                        db.session.flush()

                    if existing_parent not in student.parents:
                        student.parents.append(existing_parent)

            student_lookup[key] = student
            results['students_created'] += 1

        # Ensure financial account exists
        if not student.financial_account:
            fin_account = StudentFinancialAccount(student_id=student.id)
            db.session.add(fin_account)
            db.session.flush()

        # 4. Create or update Subscription
        if import_subscriptions and has_plan:
            plan_name = s_info.get('plan_name')
            cycle = s_info.get('cycle', 'Monthly')
            start_date = parse_procare_date(s_info.get('start_date')) or today
            end_date = parse_procare_date(s_info.get('end_date'))
            plan_status = s_info.get('plan_status', 'Active')
            items_json = s_info.get('items', [])

            # Calculate next invoice date
            invoice_gen_day = 1
            due_day = 15
            cycle_clean = cycle.lower().replace('-', '').replace(' ', '')

            if cycle_clean == 'weekly':
                next_invoice = start_date + relativedelta(weeks=1)
            elif cycle_clean == 'biweekly':
                next_invoice = start_date + relativedelta(weeks=2)
            elif cycle_clean == 'quarterly':
                try:
                    next_invoice = start_date.replace(day=invoice_gen_day)
                except ValueError:
                    next_invoice = start_date
                if start_date.day > invoice_gen_day:
                    next_invoice += relativedelta(months=3)
            else:  # Monthly
                try:
                    next_invoice = start_date.replace(day=invoice_gen_day)
                except ValueError:
                    next_invoice = start_date
                if start_date.day > invoice_gen_day:
                    next_invoice += relativedelta(months=1)

            # Check existing subscription for this student and plan
            sub = Subscription.query.filter_by(
                account_id=student.financial_account.id,
                plan_name=plan_name
            ).first()

            if sub:
                sub.status = plan_status
                sub.cycle = cycle
                sub.start_date = start_date
                sub.end_date = end_date
                sub.items_json = items_json
                sub.next_invoice_date = next_invoice
                results['subscriptions_updated'] += 1
            else:
                sub = Subscription(
                    account_id=student.financial_account.id,
                    plan_name=plan_name,
                    status=plan_status,
                    cycle=cycle,
                    start_date=start_date,
                    end_date=end_date,
                    invoice_generation_day=invoice_gen_day,
                    due_day=due_day,
                    next_invoice_date=next_invoice,
                    items_json=items_json
                )
                db.session.add(sub)
                results['subscriptions_created'] += 1

    if actor:
        log_activity(
            actor,
            f"Imported Procare plans: {results['templates_created'] + results['templates_updated']} template(s), "
            f"{results['students_created']} student(s) created, {results['subscriptions_created'] + results['subscriptions_updated']} subscription(s) active."
        )

    db.session.commit()
    return results

